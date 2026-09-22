"""ทดสอบการ export Excel: มีครบทุกชีต คอลัมน์ และอ่านกลับด้วย pandas ได้"""

from __future__ import annotations

import io

import pandas as pd
from openpyxl import load_workbook

from datetime import date

import pytest

from omr.exporter import (
    DETAIL_COLUMNS, KEY_COLUMNS, REPORT_COLUMNS, REPORT_HEADER_ROW, SHEET_DETAIL, SHEET_KEY, SHEET_REPORT,
    SHEET_STATS, SHEET_SUMMARY, SHEET_UNMATCHED, STATS_COLUMNS, SUMMARY_COLUMNS, ExamInfo, export_excel,
)
from omr.roster import RosterError, parse_roster
from omr.grader import grade
from omr.models import ReadResult, SheetResult


def _sheet(name: str, sid: str, answers: dict[int, tuple[str, str]], key: dict[int, str]) -> SheetResult:
    reads = {q: ReadResult(a, s, 0.9 if s == "OK" else 0.3, s != "OK", {}) for q, (a, s) in answers.items()}
    res = SheetResult(filename=name, student_id=sid, id_valid="?" not in sid, id_confidence=0.9)
    res.answer_reads = reads
    return grade(reads, key, res)


def test_export_excel_structure() -> None:
    key = {1: "A", 2: "B", 3: "C"}
    ok = _sheet("s1.jpg", "65001", {1: ("A", "OK"), 2: ("B", "OK"), 3: ("D", "OK")}, key)
    review = _sheet("s2.jpg", "650?1", {1: ("AB", "MULTI"), 2: ("", "BLANK"), 3: ("C", "FAINT")}, key)
    err = SheetResult(filename="bad.jpg", error="พบมาร์กเกอร์ 2 จาก 4 ตัว")
    data = export_excel([ok, review, err], key)

    xls = pd.read_excel(io.BytesIO(data), sheet_name=None)
    assert list(xls) == [SHEET_REPORT, SHEET_SUMMARY, SHEET_DETAIL, SHEET_STATS, SHEET_KEY]
    assert list(xls[SHEET_SUMMARY].columns) == SUMMARY_COLUMNS
    assert list(xls[SHEET_DETAIL].columns) == DETAIL_COLUMNS
    assert list(xls[SHEET_STATS].columns) == STATS_COLUMNS
    assert list(xls[SHEET_KEY].columns) == KEY_COLUMNS

    summary = xls[SHEET_SUMMARY]
    assert len(summary) == 3
    assert list(summary["สถานะ"]) == ["สำเร็จ", "ต้องตรวจสอบ", "ผิดพลาด"]
    assert summary.loc[0, "คะแนน"] == 2
    assert len(xls[SHEET_DETAIL]) == 6
    stats = xls[SHEET_STATS]
    assert stats.loc[0, "จำนวนตอบถูก"] == 1 and stats.loc[0, "ฝนซ้ำ"] == 1
    assert stats.loc[1, "ไม่ฝน"] == 1

    wb = load_workbook(io.BytesIO(data))
    ws = wb[SHEET_SUMMARY]
    assert ws.freeze_panes == "A2"
    zero = export_excel([_sheet("z.jpg", "01234", {1: ("A", "OK")}, key)], key)
    zws = load_workbook(io.BytesIO(zero))[SHEET_SUMMARY]
    assert zws["B2"].value == "01234" and zws["B2"].number_format == "@"
    assert ws["A1"].font.bold
    assert ws["A3"].fill.fgColor.rgb.endswith("FFF6C8")  # แถวต้องตรวจสอบ
    assert ws["A4"].fill.fgColor.rgb.endswith("F8D0D0")  # แถว error


def test_export_empty() -> None:
    xls = pd.read_excel(io.BytesIO(export_excel([], {1: "A"})), sheet_name=None)
    assert len(xls[SHEET_SUMMARY]) == 0 and len(xls[SHEET_KEY]) == 1


ROSTER_CSV = """เลขที่,รหัส,ชื่อ,นามสกุล,ห้อง
2,65002,สมหญิง,รักเรียน,ม.3/1
1,65001,สมชาย,ใจดี,ม.3/1
1,123,มานี,มีนา,ม.3/2
""".encode("utf-8")


def test_parse_roster_restores_leading_zeros_and_joins_names() -> None:
    roster = parse_roster(ROSTER_CSV, "list.csv")
    assert set(roster) == {"65001", "65002", "00123"}  # Excel ตัด 0 นำหน้าของ 00123 ทิ้ง ต้องเติมคืน
    assert roster["65001"].name == "สมชาย ใจดี" and roster["65001"].room == "ม.3/1"
    assert roster["00123"].number == 1


def test_parse_roster_from_excel_with_single_name_column() -> None:
    buf = io.BytesIO()
    pd.DataFrame({"Student ID": [1234.0, 65001], "ชื่อ-สกุล": ["ก ข", "ค ง"]}).to_excel(buf, index=False)
    roster = parse_roster(buf.getvalue(), "list.xlsx")
    assert roster["01234"].name == "ก ข" and roster["01234"].room == ""


@pytest.mark.parametrize("data,msg", [
    ("ชื่อ,ห้อง\nก,1\n", "รหัส"),
    ("รหัส,ชื่อ\n65001,ก\n65001,ข\n", "ซ้ำ"),
])
def test_parse_roster_errors(data: str, msg: str) -> None:
    with pytest.raises(RosterError, match=msg):
        parse_roster(data.encode("utf-8"), "bad.csv")


def test_report_sheet_per_room_with_names_absent_and_unmatched() -> None:
    key = {1: "A", 2: "B", 3: "C"}
    somchai = _sheet("a.jpg", "65001", {1: ("A", "OK"), 2: ("B", "OK"), 3: ("C", "OK")}, key)
    manee = _sheet("b.jpg", "00123", {1: ("A", "OK"), 2: ("C", "OK"), 3: ("C", "OK")}, key)
    stranger = _sheet("c.jpg", "99999", {1: ("A", "OK"), 2: ("B", "OK"), 3: ("D", "OK")}, key)
    exam = ExamInfo(title="กลางภาค คณิตศาสตร์", exam_date=date(2026, 9, 22))
    data = export_excel([somchai, manee, stranger], key, parse_roster(ROSTER_CSV, "list.csv"), exam)

    wb = load_workbook(io.BytesIO(data))
    # ห้อง ม.3/1 ใช้ / ในชื่อชีตไม่ได้ จึงเป็น ม.3-1
    assert wb.sheetnames[:3] == ["ม.3-1", "ม.3-2", SHEET_UNMATCHED]
    ws = wb["ม.3-1"]
    assert "กลางภาค คณิตศาสตร์" in ws["A1"].value
    assert "ม.3/1" in ws["A2"].value and "22/09/2026" in ws["A2"].value and "คะแนนเต็ม: 3" in ws["A2"].value
    assert [c.value for c in ws[REPORT_HEADER_ROW]] == REPORT_COLUMNS
    first, second = ws[REPORT_HEADER_ROW + 1], ws[REPORT_HEADER_ROW + 2]
    # เรียงตามเลขที่: สมชาย (เลขที่ 1, ได้ 3 คะแนน) แล้วสมหญิง (เลขที่ 2, ไม่ส่ง)
    assert [c.value for c in first[:5]] == [1, "65001", "สมชาย ใจดี", 3, 3]
    assert second[2].value == "สมหญิง รักเรียน" and second[3].value is None and "ขาดสอบ" in second[6].value
    assert wb["ม.3-2"][REPORT_HEADER_ROW + 1][3].value == 2
    unmatched = pd.read_excel(io.BytesIO(data), sheet_name=SHEET_UNMATCHED, dtype={"รหัสที่อ่านได้": str})
    assert list(unmatched["รหัสที่อ่านได้"]) == ["99999"] and "ไม่มีรหัสนี้" in unmatched.loc[0, "สาเหตุ"]
    summary = pd.read_excel(io.BytesIO(data), sheet_name=SHEET_SUMMARY)
    assert list(summary["ชื่อ-สกุล"].fillna("")) == ["สมชาย ใจดี", "มานี มีนา", ""]


def test_report_flags_two_sheets_with_same_student_id() -> None:
    key = {1: "A"}
    a = _sheet("a.jpg", "65001", {1: ("A", "OK")}, key)
    b = _sheet("b.jpg", "65001", {1: ("B", "OK")}, key)
    wb = load_workbook(io.BytesIO(export_excel([a, b], key, parse_roster(ROSTER_CSV, "list.csv"))))
    note = wb["ม.3-1"][REPORT_HEADER_ROW + 1][6].value
    assert "2 แผ่น" in note and "a.jpg" in note and "b.jpg" in note
