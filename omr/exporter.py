"""ส่งออกผลการตรวจเป็นไฟล์ Excel (openpyxl)"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from omr import config as C
from omr.grader import AnswerKey
from omr.models import STATUS_BLANK, STATUS_FAINT, STATUS_MULTI, STATUS_OK, SheetResult
from omr.roster import Roster, Student, rooms

log = logging.getLogger(__name__)

SHEET_REPORT = "รายงานคะแนน"          # ใช้เมื่อไม่มีรายชื่อ หรือรายชื่อไม่ได้แบ่งห้อง
SHEET_UNMATCHED = "ตรวจสอบรหัส"        # กระดาษคำตอบที่หาชื่อในรายชื่อไม่เจอ
SHEET_SUMMARY = "สรุปคะแนน"
SHEET_DETAIL = "รายข้อ"
SHEET_STATS = "สถิติรายข้อ"
SHEET_KEY = "เฉลย"

REPORT_COLUMNS = ["เลขที่", "รหัสนักศึกษา", "ชื่อ-สกุล", "คะแนน", "เต็ม", "ร้อยละ", "หมายเหตุ"]
UNMATCHED_COLUMNS = ["ลำดับ", "รหัสที่อ่านได้", "ชื่อไฟล์", "คะแนน", "เต็ม", "สาเหตุ"]
SUMMARY_COLUMNS = ["ลำดับ", "รหัสนักศึกษา", "ชื่อ-สกุล", "ห้อง", "ชื่อไฟล์", "คะแนน", "เต็ม", "ร้อยละ", "ความมั่นใจเฉลี่ย",
                   "ข้อไม่มั่นใจ", "ข้อไม่ฝน", "ข้อฝนซ้ำ", "สถานะ", "หมายเหตุ"]
DETAIL_COLUMNS = ["รหัสนักศึกษา", "ข้อ", "คำตอบ", "เฉลย", "ผล", "สถานะ", "ความมั่นใจ(%)", "ปัญหาที่ตรวจพบ"]
STATS_COLUMNS = ["ข้อ", "เฉลย", "จำนวนตอบถูก", "ร้อยละตอบถูก", "จำนวน A", "จำนวน B", "จำนวน C", "จำนวน D",
                 "ไม่ฝน", "ฝนซ้ำ"]
KEY_COLUMNS = ["ข้อ", "เฉลย"]
TEXT_COLUMNS = {"รหัสนักศึกษา", "รหัสที่อ่านได้"}

STATUS_TH = {STATUS_OK: "ปกติ", STATUS_FAINT: "ฝนจาง", STATUS_MULTI: "ฝนซ้ำ", STATUS_BLANK: "ไม่ฝน"}

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
REVIEW_FILL = PatternFill("solid", fgColor="FFF6C8")
ERROR_FILL = PatternFill("solid", fgColor="F8D0D0")
ABSENT_FILL = PatternFill("solid", fgColor="E5E7EB")  # นักเรียนในรายชื่อที่ไม่มีกระดาษคำตอบ
TITLE_FONT = Font(bold=True, size=14)
INFO_FONT = Font(bold=True)
REPORT_HEADER_ROW = 5  # แถวหัวตารางของชีตรายงาน (แถว 1–3 เป็นหัวรายงาน แถว 4 เว้นว่าง)
SHEET_TITLE_BAD_CHARS = "/\\?*[]:"  # อักขระที่ Excel ห้ามใช้ในชื่อชีต
MAX_COL_WIDTH = 60
MIN_COL_WIDTH = 8
WIDTH_PADDING = 3


@dataclass(frozen=True)
class ExamInfo:
    """ข้อมูลหัวรายงาน: การสอบอะไร ห้องไหน วันไหน"""

    title: str = ""
    room: str = ""  # ใช้เมื่อรายชื่อไม่ได้ระบุห้อง (หรือไม่มีรายชื่อ)
    exam_date: date | None = None


def _write_table(ws: Worksheet, columns: list[str], rows: list[list[Any]], fills: list[PatternFill | None],
                 header_row: int = 1) -> None:
    """เขียนตารางพร้อมหัวตาราง freeze และปรับความกว้างคอลัมน์ (header_row > 1 = มีหัวรายงานอยู่ด้านบน)"""
    for idx, col in enumerate(columns, start=1):
        ws.cell(row=header_row, column=idx, value=col)
    for cell in ws[header_row]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    text_cols = [i for i, c in enumerate(columns) if c in TEXT_COLUMNS]
    for n, (row, fill) in enumerate(zip(rows, fills), start=header_row + 1):
        for idx, value in enumerate(row, start=1):
            ws.cell(row=n, column=idx, value=value)
        cells = ws[n]
        for i in text_cols:
            cells[i].number_format = "@"  # รหัสขึ้นต้นด้วย 0 ต้องเป็นข้อความ
        if fill is not None:
            for cell in cells:
                cell.fill = fill
    ws.freeze_panes = f"A{header_row + 1}"
    for idx, col in enumerate(columns, start=1):
        values = [col] + [r[idx - 1] for r in rows]
        # ตัวอักษรไทยแคบกว่าละติน ใช้ความยาวสตริงเป็นค่าประมาณ
        width = max(len(str(v)) if v is not None else 0 for v in values) + WIDTH_PADDING
        ws.column_dimensions[get_column_letter(idx)].width = max(MIN_COL_WIDTH, min(MAX_COL_WIDTH, width))


def _note(r: SheetResult) -> str:
    notes: list[str] = []
    if r.error:
        notes.append(r.error)
    else:
        if not r.id_valid:
            notes.append("อ่านรหัสนักศึกษาไม่ได้บางหลัก")
        if r.n_uncertain:
            notes.append(f"ไม่มั่นใจ {r.n_uncertain} ข้อ")
        if r.n_multi:
            notes.append(f"ฝนซ้ำ {r.n_multi} ข้อ")
        if r.rotated:
            notes.append("ภาพกลับหัว (หมุนอัตโนมัติ)")
        notes.extend(r.warnings)
        n_ai = sum(b.source == "yolo" for b in r.problem_boxes)
        if n_ai:
            notes.append(f"YOLOv8n พบบริเวณที่มีปัญหา {n_ai} จุด")
        if r.reviewed:
            notes.append("แก้ไขโดยผู้ตรวจ")
    return "; ".join(notes)


def _sheet_title(room: str, used: set[str]) -> str:
    """ชื่อชีตห้ามมี / \\ ? * [ ] : ยาวไม่เกิน 31 ตัว และห้ามซ้ำ (ห้อง ม.3/1 จึงเป็นชีต ม.3-1)"""
    for ch in SHEET_TITLE_BAD_CHARS:
        room = room.replace(ch, "-")
    base = (room.strip() or SHEET_REPORT)[:28]
    title, n = base, 2
    while title in used:
        title, n = f"{base} ({n})", n + 1
    used.add(title)
    return title


def _score_cells(r: SheetResult) -> list[Any]:
    return [r.score, r.total, r.percent] if not r.error else [None, None, None]


def _report_note(r: SheetResult) -> str:
    if r.error:
        return "ตรวจไม่ได้: " + r.error
    if r.needs_review:
        return "ต้องตรวจสอบ (" + _note(r) + ")"
    return "แก้ไขโดยผู้ตรวจ" if r.reviewed else ""


def _write_report(ws: Worksheet, room: str, exam: ExamInfo, total: int,
                  rows: list[list[Any]], fills: list[PatternFill | None]) -> None:
    """ชีตรายงานคะแนนสำหรับส่งต่อ: หัวรายงาน (การสอบ / ห้อง / วันที่ / สถิติ) + ตารางรายคน"""
    scores = [row[3] for row in rows if isinstance(row[3], (int, float))]
    ws["A1"] = "รายงานผลการสอบ" + (f" — {exam.title}" if exam.title else "")
    ws["A1"].font = TITLE_FONT
    exam_date = f"{exam.exam_date:%d/%m/%Y}" if exam.exam_date else "-"
    ws["A2"] = f"ห้อง: {room or '-'}      วันที่สอบ: {exam_date}      คะแนนเต็ม: {total}"
    stats = f"นักเรียน {len(rows)} คน · มีคะแนน {len(scores)} คน"
    if scores:
        stats += f" · เฉลี่ย {sum(scores) / len(scores):.2f} · สูงสุด {max(scores)} · ต่ำสุด {min(scores)}"
    ws["A3"] = stats
    ws["A2"].font = ws["A3"].font = INFO_FONT
    _write_table(ws, REPORT_COLUMNS, rows, fills, header_row=REPORT_HEADER_ROW)


def _add_report_sheets(wb: Workbook, results: list[SheetResult], roster: Roster | None, exam: ExamInfo,
                       total: int) -> None:
    """ชีตรายงานคะแนนรายห้อง (อยู่หน้าสุดของไฟล์) — จับคู่รหัสบนกระดาษคำตอบกับรายชื่อนักเรียน"""
    used: set[str] = {SHEET_UNMATCHED, SHEET_SUMMARY, SHEET_DETAIL, SHEET_STATS, SHEET_KEY}
    if not roster:
        ordered = sorted(results, key=lambda r: (r.error is not None, r.student_id, r.filename))
        rows = [[i, r.student_id, "", *_score_cells(r), _report_note(r)] for i, r in enumerate(ordered, start=1)]
        fills = [ERROR_FILL if r.error else REVIEW_FILL if r.needs_review else None for r in ordered]
        _write_report(wb.create_sheet(_sheet_title(exam.room, used)), exam.room, exam, total, rows, fills)
        return

    by_id: dict[str, list[SheetResult]] = {}
    for r in results:
        if not r.error and r.id_valid:
            by_id.setdefault(r.student_id, []).append(r)
    for room in rooms(roster):
        students = sorted((s for s in roster.values() if s.room == room),
                          key=lambda s: (s.number is None, s.number or 0, s.student_id))
        rows, fills = [], []
        for s in students:
            sheets = by_id.get(s.student_id, [])
            if not sheets:
                rows.append([s.number, s.student_id, s.name, None, None, None, "ไม่พบกระดาษคำตอบ (ขาดสอบ?)"])
                fills.append(ABSENT_FILL)
                continue
            r = sheets[0]
            note = _report_note(r)
            if len(sheets) > 1:
                files = ", ".join(x.filename for x in sheets)
                note = f"พบกระดาษคำตอบ {len(sheets)} แผ่นใช้รหัสนี้ ({files}) — ตรวจสอบ" + (f"; {note}" if note else "")
            rows.append([s.number, s.student_id, s.name, *_score_cells(r), note])
            fills.append(REVIEW_FILL if len(sheets) > 1 or r.needs_review else None)
        label = room or exam.room
        _write_report(wb.create_sheet(_sheet_title(label, used)), label, exam, total, rows, fills)

    # กระดาษคำตอบที่ไม่รู้ว่าเป็นของใคร ต้องให้ครูเห็น ไม่ใช่หายไปเงียบ ๆ
    unmatched = [r for r in results if r.error or not r.id_valid or r.student_id not in roster]
    if unmatched:
        rows = []
        for i, r in enumerate(unmatched, start=1):
            why = ("ตรวจไม่ได้: " + r.error if r.error else "อ่านรหัสไม่ครบทุกหลัก" if not r.id_valid
                   else "ไม่มีรหัสนี้ในรายชื่อนักเรียน")
            rows.append([i, r.student_id, r.filename, *_score_cells(r)[:2], why])
        _write_table(wb.create_sheet(SHEET_UNMATCHED), UNMATCHED_COLUMNS, rows, [ERROR_FILL] * len(rows))


def export_excel(results: list[SheetResult], answer_key: AnswerKey, roster: Roster | None = None,
                 exam: ExamInfo | None = None) -> bytes:
    """สร้างไฟล์ Excel คืนค่าเป็น bytes

    ชีตแรกคือรายงานคะแนน (แยกชีตละห้องเมื่อรายชื่อมีหลายห้อง) ตามด้วยชีตสรุป / รายข้อ / สถิติ / เฉลย
    """
    exam = exam or ExamInfo()
    wb = Workbook()
    wb.remove(wb.active)
    _add_report_sheets(wb, results, roster, exam, len(answer_key))
    ws_sum = wb.create_sheet(SHEET_SUMMARY)

    rows, fills = [], []
    for i, r in enumerate(results, start=1):
        student: Student | None = (roster or {}).get(r.student_id) if r.id_valid else None
        rows.append([
            i, r.student_id, student.name if student else "", student.room if student else "", r.filename, r.score if not r.error else None, r.total if not r.error else None,
            r.percent if not r.error else None, round(r.avg_confidence * 100, 1) if not r.error else None,
            r.n_uncertain, r.n_blank, r.n_multi, r.status_label, _note(r),
        ])
        fills.append(ERROR_FILL if r.error else REVIEW_FILL if r.needs_review else None)
    _write_table(ws_sum, SUMMARY_COLUMNS, rows, fills)

    ws_det = wb.create_sheet(SHEET_DETAIL)
    rows, fills = [], []
    for r in results:
        for q in r.questions:
            rows.append([
                r.student_id, q.number, q.answer, q.key, "ถูก" if q.is_correct else "ผิด",
                STATUS_TH.get(q.status, q.status) + (" (แก้ไข)" if q.edited else ""), round(q.confidence * 100, 1),
                q.problem,
            ])
            flagged = q.uncertain or q.status in (STATUS_MULTI, STATUS_BLANK)
            fills.append(REVIEW_FILL if flagged else None)
    _write_table(ws_det, DETAIL_COLUMNS, rows, fills)

    ws_stat = wb.create_sheet(SHEET_STATS)
    graded = [r for r in results if not r.error]
    rows = []
    for qn, key in sorted(answer_key.items()):
        answers = [q for r in graded for q in r.questions if q.number == qn]
        n = len(answers)
        correct = sum(q.is_correct for q in answers)
        counts = [sum(q.answer == ch and q.status in (STATUS_OK, STATUS_FAINT) for q in answers) for ch in C.CHOICES]
        rows.append([
            qn, key, correct, round(100.0 * correct / n, 1) if n else 0.0, *counts,
            sum(q.status == STATUS_BLANK for q in answers), sum(q.status == STATUS_MULTI for q in answers),
        ])
    _write_table(ws_stat, STATS_COLUMNS, rows, [None] * len(rows))

    ws_key = wb.create_sheet(SHEET_KEY)
    rows = [[q, a] for q, a in sorted(answer_key.items())]
    _write_table(ws_key, KEY_COLUMNS, rows, [None] * len(rows))

    buf = io.BytesIO()
    wb.save(buf)
    log.info("export Excel: %d แผ่น", len(results))
    return buf.getvalue()
